import re as _re

# =============================================================================
# ACTIVITY LABELS AND COLORS
# =============================================================================

ACTIVITY_TYPE_LABELS = {
    "InvokePipeline":   "Invocar canalizacion",
    "ExecutePipeline":  "Invocar canalizacion",
    "Copy":             "Copiar datos",
    "Script":           "Script",
    "TridentNotebook":  "Bloc de notas",
    "Notebook":         "Bloc de notas",
    "RefreshDataflow":  "Dataflow",
    "ForEach":          "Para cada uno",
    "IfCondition":      "Si condicion",
    "Wait":             "Esperar",
    "SetVariable":      "Establecer variable",
    "Lookup":           "Busqueda",
    "GetMetadata":      "Obtener metadatos",
    "WebActivity":      "Web",
}

ACTIVITY_COLORS = {
    "InvokePipeline":   ((233, 30, 140),  (253, 240, 248)),
    "ExecutePipeline":  ((233, 30, 140),  (253, 240, 248)),
    "Copy":             ((59,  130, 246), (239, 246, 255)),
    "TridentNotebook":  ((124, 58,  237), (245, 243, 255)),
    "Notebook":         ((124, 58,  237), (245, 243, 255)),
    "RefreshDataflow":  ((16,  185, 129), (236, 253, 245)),
    "Script":           ((16,  185, 129), (236, 253, 245)),
    "ForEach":          ((245, 158, 11),  (255, 251, 235)),
    "IfCondition":      ((239, 68,  68),  (254, 242, 242)),
    "default":          ((96,  96,  96),  (245, 245, 245)),
}


# =============================================================================
# PIPELINE DIAGRAM
# =============================================================================

def make_pipeline_diagram_png(pipeline_name, activities, pipeline_id_to_name=None):
    from PIL import Image, ImageDraw, ImageFont
    import io as _io
    if not activities:
        return None
    if pipeline_id_to_name is None:
        pipeline_id_to_name = {}

    C_BG    = (255, 255, 255)
    C_BLACK = (30, 30, 30)

    act_map  = {a["name"]: a for a in activities}
    in_deg   = {a["name"]: 0 for a in activities}
    children = {a["name"]: [] for a in activities}
    for act in activities:
        for d in act.get("dependsOn", []):
            dep = d["activity"]
            if dep in children:
                children[dep].append(act["name"])
                in_deg[act["name"]] += 1

    from collections import deque
    queue = deque([n for n, d in in_deg.items() if d == 0])
    level = {n: 0 for n in queue}
    order = []
    while queue:
        n = queue.popleft()
        order.append(n)
        for ch in children[n]:
            in_deg[ch] -= 1
            level[ch] = max(level.get(ch, 0), level[n] + 1)
            if in_deg[ch] == 0:
                queue.append(ch)

    max_level = max(level.values()) if level else 0
    by_level  = {}
    for n, l in level.items():
        by_level.setdefault(l, []).append(n)

    BOX_W = 130; BOX_H = 54
    H_GAP = 20;  V_GAP = 14
    PAD   = 30;  SCALE = 2

    max_col_h = max(len(v) for v in by_level.values()) if by_level else 1
    W = PAD*2 + (max_level+1)*BOX_W + max_level*H_GAP
    H = PAD*2 + max_col_h*BOX_H + (max_col_h-1)*V_GAP
    W = max(W, 300); H = max(H, 120)

    img  = Image.new("RGB", (W*SCALE, H*SCALE), C_BG)
    draw = ImageDraw.Draw(img)

    def s(v): return v * SCALE

    def get_font(size, bold=False):
        try:
            return ImageFont.truetype(
                "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
                if bold else
                "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", s(size))
        except Exception:
            return ImageFont.load_default()

    positions = {}
    for lvl in range(max_level + 1):
        acts_at  = by_level.get(lvl, [])
        n_acts   = len(acts_at)
        total_h  = n_acts * BOX_H + (n_acts-1) * V_GAP
        start_y  = (H - total_h) // 2
        x        = PAD + lvl * (BOX_W + H_GAP)
        for i, aname in enumerate(acts_at):
            positions[aname] = (x, start_y + i*(BOX_H+V_GAP))

    for act in activities:
        if act["name"] not in positions:
            continue
        x2, y2 = positions[act["name"]]
        cy2 = y2 + BOX_H//2
        for dep in act.get("dependsOn", []):
            dn   = dep["activity"]
            cond = dep.get("dependencyConditions", [""])[0]
            if dn not in positions:
                continue
            x1, y1 = positions[dn]
            cx1 = x1 + BOX_W; cy1 = y1 + BOX_H//2
            draw.line([(s(cx1),s(cy1)),(s(x2),s(cy2))], fill=(200,200,200), width=s(2))
            draw.polygon([(s(x2+1),s(cy2)),(s(x2-7),s(cy2-4)),(s(x2-7),s(cy2+4))],
                         fill=(180,180,180))
            if cond:
                draw.text((s((cx1+x2)//2), s((cy1+cy2)//2-8)), cond,
                          fill=(150,150,150), font=get_font(5), anchor="mm")

    for act in activities:
        aname = act["name"]
        atype = act.get("type", "default")
        if aname not in positions:
            continue
        x, y = positions[aname]
        oc, fc = ACTIVITY_COLORS.get(atype, ACTIVITY_COLORS["default"])
        draw.rounded_rectangle([s(x+2),s(y+2),s(x+BOX_W+2),s(y+BOX_H+2)],
                                radius=s(6), fill=(220,220,220))
        draw.rounded_rectangle([s(x),s(y),s(x+BOX_W),s(y+BOX_H)],
                                radius=s(6), fill=fc, outline=oc, width=s(2))
        tl = ACTIVITY_TYPE_LABELS.get(atype, atype)
        draw.rounded_rectangle([s(x),s(y),s(x+BOX_W),s(y+18)],
                                radius=s(6), fill=oc, outline=oc, width=s(1))
        draw.rectangle([s(x),s(y+12),s(x+BOX_W),s(y+18)], fill=oc)
        draw.text((s(x+BOX_W//2),s(y+9)), tl, fill=(255,255,255),
                  font=get_font(6,bold=True), anchor="mm")
        dn = aname
        if atype in ("InvokePipeline","ExecutePipeline"):
            pid = act.get("typeProperties",{}).get("pipelineId","")
            if pid in pipeline_id_to_name:
                dn = pipeline_id_to_name[pid]
        words = dn.replace("_"," ").split()
        mid   = max(1, len(words)//2)
        l1    = " ".join(words[:mid])
        l2    = " ".join(words[mid:])
        fn    = get_font(7, bold=True)
        if l2:
            draw.text((s(x+BOX_W//2),s(y+30)), l1, fill=C_BLACK, font=fn, anchor="mm")
            draw.text((s(x+BOX_W//2),s(y+42)), l2, fill=C_BLACK, font=fn, anchor="mm")
        else:
            draw.text((s(x+BOX_W//2),s(y+36)), l1, fill=C_BLACK, font=fn, anchor="mm")

    img = img.resize((W,H), Image.LANCZOS)
    buf = _io.BytesIO()
    img.save(buf, format="PNG", dpi=(150,150))
    return buf.getvalue()


def get_pipeline_diagram_path(pipeline_name, activities,
                               pipeline_id_to_name=None, suffix="main"):
    import tempfile, os
    png = make_pipeline_diagram_png(pipeline_name, activities, pipeline_id_to_name)
    if not png:
        return None
    path = os.path.join(tempfile.gettempdir(), f"pipeline_{suffix}.png")
    with open(path, "wb") as f:
        f.write(png)
    print(f"  OK Diagrama pipeline '{pipeline_name}'")
    return path


# =============================================================================
# TECHNICAL TEXT GENERATORS
# =============================================================================

def _generate_main_pipeline_text(pipeline_name, activities, schedule_time, pid_to_name):
    if not activities:
        return (f"La canalizacion {pipeline_name} gestiona el flujo de "
                f"actualizacion programado diariamente a las {schedule_time}.")
    type_counts  = {}
    invoke_names = []
    has_endpoint = False
    conditions   = set()
    dependencies = {}
    for act in activities:
        atype = act.get("type","")
        aname = act.get("name","")
        type_counts[atype] = type_counts.get(atype,0) + 1
        dependencies[aname] = [d["activity"] for d in act.get("dependsOn",[])]
        for d in act.get("dependsOn",[]):
            cond = d.get("dependencyConditions",[""])[0]
            if cond: conditions.add(cond)
        if atype in ("InvokePipeline","ExecutePipeline"):
            pid   = act.get("typeProperties",{}).get("pipelineId","")
            invoke_names.append(pid_to_name.get(pid, aname))
        if "endpoint" in aname.lower() or "sql" in aname.lower():
            has_endpoint = True
    roots = [n for n,d in dependencies.items() if not d]
    sequential = len(roots) <= 1
    layers = []
    for n in invoke_names + list(dependencies.keys()):
        nl = n.lower()
        if "bronze" in nl and "Bronze" not in layers: layers.append("Bronze")
        if "silver" in nl and "Silver" not in layers: layers.append("Silver")
        if "gold"   in nl and "Gold"   not in layers: layers.append("Gold")
    parts = []
    n_acts = len(activities)
    if layers:
        ls = " -> ".join(layers)
        parts.append(
            f"La canalizacion **{pipeline_name}** implementa un patron de "
            f"orquestacion {'secuencial' if sequential else 'con paralelismo parcial'} "
            f"en {len(layers)} capas ({ls}), compuesta por {n_acts} actividades "
            f"y programada diariamente a las **{schedule_time}**.")
    else:
        parts.append(
            f"La canalizacion **{pipeline_name}** esta compuesta por {n_acts} "
            f"actividades y programada diariamente a las **{schedule_time}**.")
    if conditions:
        parts.append(
            f"La integridad del flujo se garantiza mediante condiciones de "
            f"dependencia ({', '.join(sorted(conditions))}), asegurando que "
            f"cada capa se procese solo cuando la anterior ha concluido.")
    if has_endpoint:
        ep = next((a["name"] for a in activities
                   if "endpoint" in a["name"].lower() or "sql" in a["name"].lower()), None)
        if ep:
            parts.append(
                f"La actividad **{ep}** actua como barrera de sincronizacion, "
                f"asegurando que el endpoint SQL Analytics este disponible antes "
                f"de iniciar las transformaciones de capas superiores.")
    if not sequential and len(roots) > 1:
        rs = ", ".join(f"**{r}**" for r in roots[:3])
        parts.append(
            f"Las actividades {rs} se ejecutan en paralelo al inicio del flujo, "
            f"optimizando el tiempo total de procesamiento.")
    return " ".join(parts)


def _generate_layer_pipeline_text(pipeline_name, activities, layer):
    if not activities:
        return f"La canalizacion **{pipeline_name}** gestiona la capa {layer}."
    n_acts       = len(activities)
    type_counts  = {}
    has_notebook = False
    has_dataflow = False
    pre_copy_nb  = None
    parallel_count = sum(1 for a in activities if not a.get("dependsOn"))
    for act in activities:
        atype = act.get("type","")
        aname = act.get("name","")
        type_counts[atype] = type_counts.get(atype,0) + 1
        if atype in ("TridentNotebook","Notebook"):
            has_notebook = True
            if "precopy" in aname.lower() or "pre_copy" in aname.lower():
                pre_copy_nb = aname
        if atype == "RefreshDataflow":
            has_dataflow = True
    type_labels = []
    for atype, cnt in type_counts.items():
        type_labels.append(f"{cnt} de tipo {ACTIVITY_TYPE_LABELS.get(atype,atype)}")
    types_str = ", ".join(type_labels)
    parts = []
    if layer.lower() == "bronze":
        parts.append(
            f"La canalizacion **{pipeline_name}** gestiona la ingesta de datos "
            f"en bruto desde los sistemas de origen. Esta compuesta por {n_acts} "
            f"actividades ({types_str}).")
        if parallel_count > 1:
            parts.append(
                f"Las {parallel_count} actividades principales se ejecutan en "
                f"paralelo, optimizando el tiempo de carga inicial.")
        if has_notebook:
            parts.append(
                "Incorpora un notebook PySpark para la carga completa o incremental "
                "de las tablas segun la configuracion definida en los ficheros de parametros.")
        if has_dataflow:
            parts.append(
                "Un Dataflow gestiona la ingesta de ficheros Excel desde el sitio "
                "SharePoint, almacenandolos en el esquema Bronze del Lakehouse.")
        parts.append(
            "Su objetivo es realizar una copia fiel de los datos de origen sin "
            "transformaciones, garantizando la trazabilidad hacia las capas superiores.")
    elif layer.lower() == "silver":
        parts.append(
            f"La canalizacion **{pipeline_name}** aplica transformaciones intermedias "
            f"sobre los datos Bronze. Sus {n_acts} actividades ({types_str}) generan "
            f"las tablas de la capa Silver.")
        if parallel_count > 1:
            parts.append(
                f"Las {parallel_count} transformaciones se ejecutan en paralelo, "
                f"reduciendo el tiempo total de procesamiento.")
        parts.append(
            "Estas tablas actuan como puente hacia el modelo analitico final, "
            "aplicando reglas de limpieza y controles de calidad sobre los datos en bruto.")
    elif layer.lower() == "gold":
        n_td = sum(1 for a in activities if "_TD_" in a.get("name","").upper())
        n_th = sum(1 for a in activities if "_TH_" in a.get("name","").upper())
        parts.append(
            f"La canalizacion **{pipeline_name}** construye el modelo dimensional "
            f"final compuesto por {n_acts} actividades ({types_str}).")
        if pre_copy_nb:
            parts.append(
                f"La actividad **{pre_copy_nb}** realiza un borrado parcial previo "
                f"al volcado de datos, garantizando la consistencia del modelo.")
        if n_td > 0 or n_th > 0:
            parts.append(
                f"Se cargan {n_td} tablas de dimension y {n_th} tablas de hechos "
                f"respetando las dependencias del modelo.")
    return " ".join(parts)


def _add_technical_text(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after  = Pt(4)
    p.paragraph_format.line_spacing = Pt(13)
    parts = _re.split(r'(\*\*[^*]+\*\*)', text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            run.font.name = "Arial"; run.font.size = Pt(9); run.bold = True
        else:
            run = p.add_run(part)
            run.font.name = "Arial"; run.font.size = Pt(9)
    return p


# =============================================================================
# SECTION 4 -- ORCHESTRATION
# =============================================================================

def detect_main_pipeline_for_layer(all_pipelines, layer_name):
    layer_lower = layer_name.lower()
    candidates  = [
        p for p in all_pipelines
        if layer_lower in p.get("name","").lower()
        and not any(kw in p.get("name","").lower()
                    for kw in ["carga", "000", "complete", "master", "main"])
    ]
    if candidates:
        best = min(candidates, key=lambda p: len(p.get("name","")))
        print(f"  Pipeline para {layer_name}: '{best['name']}'")
        return best
    print(f"  warning Pipeline para capa {layer_name} no detectado")
    return None


def _build_layer_subsection(doc, heading, layer_name, folder_name,
                             pipeline_name, activities, pid_to_name,
                             figure_n, ws_id=None, all_pipelines=None):
    add_h2(doc, heading)
    add_body(doc,
        f"La canalizacion principal de {layer_name} "
        f"\"{folder_name}\" tiene el siguiente formato:")
    doc.add_paragraph()
    acts   = activities
    p_name = pipeline_name
    if not acts and ws_id and all_pipelines:
        p = detect_main_pipeline_for_layer(all_pipelines, layer_name)
        if p:
            p_name = p["name"]
            acts   = get_pipeline_definition(ws_id, p["id"])
    diag_path = get_pipeline_diagram_path(p_name, acts, pid_to_name,
                                           suffix=layer_name.lower())
    insert_figure(doc, figure_n, diag_path,
        f"Figura {figure_n}. Actividades que componen el pipeline "
        f"de carga de la capa {layer_name}.",
        width_inches=5.5)
    doc.add_paragraph()
    tech_text = _generate_layer_pipeline_text(p_name, acts, layer_name)
    _add_technical_text(doc, tech_text)
    doc.add_paragraph()


def build_orchestration_section(doc, ws_list, folders, all_pipelines):
    add_h1(doc, "4. Capa de orquestacion")
    add_body(doc, TEXT_ORQUESTACION)
    doc.add_paragraph()

    bronze_n = folders.get("bronze","001 - Data - Bronze")
    silver_n = folders.get("silver","002 - Data - Silver")
    gold_n   = folders.get("gold",  "003 - Data - Gold")

    # 4.1
    add_h2(doc, "4.1 Estructura de carpetas")
    add_body(doc, TEXT_ORQUESTACION_4_1)
    doc.add_paragraph()
    for fn, desc in [
        (bronze_n, "Incluye elementos para la carga de datos sin procesar."),
        (silver_n, "Incluye elementos para la transformacion y limpieza de datos."),
        (gold_n,   "Incluye elementos para generar las tablas finales y el modelo."),
    ]:
        p  = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        rb = p.add_run(fn);   rb.font.name="Arial"; rb.font.size=Pt(9); rb.bold=True
        rs = p.add_run(f": {desc}"); rs.font.name="Arial"; rs.font.size=Pt(9)
    doc.add_paragraph()
    diag_folder = get_folder_diagram_path(ws_list[0]["name"], folders)
    insert_figure(doc, 3, diag_folder,
        "Figura 3. Estructura de carpetas segun la arquitectura de medallas",
        width_inches=5.5)
    doc.add_paragraph()

    # 4.2
    add_h2(doc, "4.2 Canalizaciones de datos (Pipelines)")
    add_body(doc, TEXT_ORQUESTACION_4_2)
    doc.add_paragraph()
    add_body(doc, TEXT_ORQUESTACION_4_2b)
    doc.add_paragraph()

    # 4.2.1
    add_h2(doc, "4.2.1 Canalizacion de carga completa")
    schedule_time   = get_pipeline_schedule_time()
    main_activities = []
    main_p_name     = ""
    pid_to_name     = {p["id"]: p["name"] for p in all_pipelines}
    ws_id           = ws_list[0]["id"] if ws_list else None

    for ws in ws_list:
        acts = ws.get("main_pipeline_activities",[])
        name = ws.get("main_pipeline_name","")
        if acts:
            main_activities = acts
            main_p_name     = name
            break

    if not main_activities and all_pipelines and ws_id:
        main_p = detect_main_pipeline(all_pipelines)
        if main_p:
            main_p_name     = main_p["name"]
            main_activities = get_pipeline_definition(ws_id, main_p["id"])

    add_body(doc,
        f"Esta es la canalizacion principal, se encarga de actualizar todos "
        f"los datos de la infraestructura. Esta programado para ejecutarse "
        f"todos los dias a las {schedule_time}.")
    doc.add_paragraph()

    diag_path = get_pipeline_diagram_path(
        main_p_name, main_activities, pid_to_name, suffix="orchestrator")
    insert_figure(doc, 4, diag_path,
        "Figura 4. Diagrama del flujo de actividades de la carga diaria "
        "para la actualizacion de las capas bronze, silver y gold.",
        width_inches=5.8)
    doc.add_paragraph()

    tech_text = _generate_main_pipeline_text(
        main_p_name, main_activities, schedule_time, pid_to_name)
    _add_technical_text(doc, tech_text)
    doc.add_paragraph()

    if main_activities:
        add_body(doc, "Actividades de la canalizacion principal:")
        doc.add_paragraph()
        rows = []
        for act in main_activities:
            atype    = act.get("type","")
            aname    = act.get("name","")
            deps     = [d["activity"] for d in act.get("dependsOn",[])]
            conds    = [d.get("dependencyConditions",[""])[0]
                        for d in act.get("dependsOn",[])]
            rows.append([
                aname,
                ACTIVITY_TYPE_LABELS.get(atype, atype),
                ", ".join(deps) if deps else "--",
                ", ".join(c for c in conds if c) or "--"
            ])
        schema_table(doc,
            ["Actividad","Tipo","Depende de","Condicion"],
            rows, col_widths=[2400,1800,2400,1200])
    doc.add_paragraph()

    # Resolver sub-pipelines dinamicamente desde el orquestador
    layer_pipelines = {}
    for act in main_activities:
        if act.get("type") in ("InvokePipeline","ExecutePipeline"):
            pid   = act.get("typeProperties",{}).get("pipelineId","")
            pname = pid_to_name.get(pid, act.get("name",""))
            nl    = pname.lower()
            if "bronze" in nl: layer_pipelines["Bronze"] = {"name":pname,"id":pid}
            elif "silver" in nl: layer_pipelines["Silver"] = {"name":pname,"id":pid}
            elif "gold"   in nl: layer_pipelines["Gold"]   = {"name":pname,"id":pid}

    for layer in ["Bronze","Silver","Gold"]:
        if layer not in layer_pipelines:
            p = detect_main_pipeline_for_layer(all_pipelines, layer)
            if p:
                layer_pipelines[layer] = {"name":p["name"],"id":p["id"]}

    layer_config = [
        ("4.2.1.1 Capa Bronze", "Bronze", bronze_n, 5),
        ("4.2.1.2 Capa Silver", "Silver", silver_n, 6),
        ("4.2.1.3 Capa Gold",   "Gold",   gold_n,   7),
    ]
    for heading, layer, folder_n, fig_n in layer_config:
        lp_info = layer_pipelines.get(layer)
        acts    = []
        pname   = ""
        if lp_info and ws_id:
            pname = lp_info["name"]
            acts  = get_pipeline_definition(ws_id, lp_info["id"])
        _build_layer_subsection(
            doc, heading, layer, folder_n,
            pname, acts, pid_to_name, fig_n,
            ws_id=ws_id, all_pipelines=all_pipelines)


def build_reports_section(doc, reports):
    if not reports:
        return
    add_h2(doc, "4.5 Reports")
    rows = sorted([
        [r.get("name",""), r.get("datasetName",""), r.get("createdBy",""),
         (r.get("modifiedDateTime","") or "")[:10]]
        for r in reports
    ], key=lambda r: (r[0],r[1]))
    schema_table(doc,
        ["Nombre","Dataset vinculado","Creado por","Modificado"],
        rows, col_widths=[2600,2600,2000,1400])


# =============================================================================
# SECTION 5 -- SEMANTIC MODELS
# =============================================================================

def build_semantic_models_section(doc, all_semantic_models):
    add_h1(doc, "5. Semantic Models / Datasets")
    if not all_semantic_models:
        add_body(doc, "No se encontraron Semantic Models.")
        return
    seen = set(); unique_sm = []
    for m in all_semantic_models:
        if m.get("id") not in seen:
            seen.add(m.get("id")); unique_sm.append(m)
    for m in sorted(unique_sm, key=lambda x: x.get("name","")):
        add_h2(doc, f"  {m.get('name','')}")
        meta_table(doc, {
            "ID":              m.get("id","--"),
            "Configurado por": m.get("configuredBy","--"),
            "Refreshable":     "Si" if m.get("isRefreshable") else "No",
        })
        rels = m.get("relationships",[])
        if rels:
            add_body(doc, "Relaciones entre tablas:")
            schema_table(doc,
                ["Tabla Origen","Col Origen","Tabla Destino","Col Destino"],
                sorted([[r.get("fromTable",""),r.get("fromColumn",""),
                         r.get("toTable",""),  r.get("toColumn","")]
                        for r in rels], key=lambda r:(r[0],r[2])))
        doc.add_paragraph()


# =============================================================================
# BACK COVER + STATIC TOC
# =============================================================================

def build_back_cover(doc):
    doc.add_page_break()
    for _ in range(10):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(_LOGO_TMP, width=Inches(2.8))
    doc.add_paragraph()
    url_p = doc.add_paragraph()
    url_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    url_run = url_p.add_run("nunsys.com")
    url_run.font.name="Arial"; url_run.font.size=Pt(11)
    url_run.font.bold=True; url_run.font.color.rgb=MAGENTA


def build_static_toc(doc):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    entries = []
    page = 1; paras = 0
    for para in doc.paragraphs:
        sn   = para.style.name if para.style else ""
        text = para.text.strip()
        for run in para.runs:
            if hasattr(run._r,"xml") and "w:br" in run._r.xml and "page" in run._r.xml:
                page += 1; paras = 0
        paras += 1
        if paras > 35: page += 1; paras = 0
        if not text: continue
        if "Heading 1" in sn: entries.append({"level":1,"text":text,"page":page})
        elif "Heading 2" in sn: entries.append({"level":2,"text":text,"page":page})
    if not entries: return doc
    toc_para = None
    for para in doc.paragraphs:
        for run in para.runs:
            if hasattr(run._r,"xml") and "TOC" in run._r.xml:
                toc_para = para; break
        if toc_para: break
    if not toc_para: return doc
    tp  = toc_para._p.getparent()
    idx = list(tp).index(toc_para._p) + 1
    for entry in reversed(entries):
        pn  = OxmlElement("w:p")
        pPr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs"); ts = OxmlElement("w:tab")
        ts.set(qn("w:val"),"right"); ts.set(qn("w:leader"),"dot"); ts.set(qn("w:pos"),"8640")
        tabs.append(ts); pPr.append(tabs)
        if entry["level"] == 2:
            ind = OxmlElement("w:ind"); ind.set(qn("w:left"),"360"); pPr.append(ind)
        pn.append(pPr)
        rt = OxmlElement("w:r"); rPr = OxmlElement("w:rPr")
        sz = OxmlElement("w:sz"); sz.set(qn("w:val"),"18" if entry["level"]==1 else "17")
        rPr.append(sz)
        if entry["level"] == 1: rPr.append(OxmlElement("w:b"))
        rt.append(rPr); te = OxmlElement("w:t")
        te.set(qn("xml:space"),"preserve"); te.text=entry["text"]; rt.append(te); pn.append(rt)
        rtab = OxmlElement("w:r"); rtab.append(OxmlElement("w:tab")); pn.append(rtab)
        rp = OxmlElement("w:r"); rPr2 = OxmlElement("w:rPr")
        sz2 = OxmlElement("w:sz"); sz2.set(qn("w:val"),"18" if entry["level"]==1 else "17")
        rPr2.append(sz2); rp.append(rPr2); tp2 = OxmlElement("w:t")
        tp2.text=str(entry["page"]); rp.append(tp2); pn.append(rp)
        tp.insert(idx, pn)
    toc_para._p.getparent().remove(toc_para._p)
    return doc


# =============================================================================
# MAIN DOCUMENT BUILDER
# =============================================================================

def build_word_document(data):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width    = Cm(21);   sec.page_height   = Cm(29.7)
    sec.top_margin    = Cm(2.5);  sec.bottom_margin = Cm(2.5)
    sec.left_margin   = Cm(2.5);  sec.right_margin  = Cm(2.5)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(9)

    build_cover(doc, data)
    build_doc_info(doc, data)

    # 1. OBJETO
    add_h1(doc, "1. Objeto")
    add_body(doc, TEXT_OBJETO.format(cliente=DOC_CLIENT))
    doc.add_paragraph()

    # 2. FABRIC WORKSPACE
    build_workspace_section(
        doc, data["workspaces"],
        data.get("all_workspaces", data["workspaces"]),
        data.get("capacity_info", CAPACITY_INFO_OVERRIDE),
        data.get("capacity_id","--"))

    # Consolidar artefactos
    ws_lakehouses_map = {}
    seen_dw  = set(); all_dw        = []
    seen_sm  = set(); all_sm        = []
    seen_p   = set(); all_pipelines = []
    seen_r   = set(); all_reports   = []
    folders  = {"bronze":"001 - Data - Bronze",
                "silver":"002 - Data - Silver",
                "gold":  "003 - Data - Gold"}
    EXCLUDE_LH = (
                "StagingLakehouseForDataflows",
                "DataflowsStagingLakehouse",
                "DataflowsStaging"
                )
    storage_folder = STORAGE_FOLDER_NAME

    for ws in data["workspaces"]:
        ws_name = ws["name"]
        for lh in ws["lakehouses"]:
            # Excluir todos los lakehouses de Staging/Dataflows
            if any(lh["name"].startswith(p) for p in EXCLUDE_LH):
                continue
            lh_name = lh["name"]
            if lh_name not in ws_lakehouses_map:
                ws_lakehouses_map[lh_name] = []
            ws_lakehouses_map[lh_name].append((ws_name, lh))
        for dw in ws["warehouses"]:
            if (not dw["name"].startswith("StagingWarehouseForDataflows")
                    and dw.get("id") not in seen_dw):
                seen_dw.add(dw.get("id")); all_dw.append(dw)
        for sm in ws["semantic_models"]:
            if sm.get("id") not in seen_sm:
                seen_sm.add(sm.get("id")); all_sm.append(sm)
        for p in ws["pipelines"]:
            if p.get("id") not in seen_p:
                seen_p.add(p.get("id")); all_pipelines.append(p)
        for r in ws["reports"]:
            if r.get("id") not in seen_r:
                seen_r.add(r.get("id")); all_reports.append(r)
        if ws.get("folder_names") and any(ws["folder_names"].values()):
            folders = ws["folder_names"]
        if ws.get("storage_folder"):
            storage_folder = ws["storage_folder"]

    # Lista plana de LH deduplicada
    seen_lh = set(); all_lh_dedup = []
    for lh_list in ws_lakehouses_map.values():
        for _, lh in lh_list:
            if lh.get("id") not in seen_lh:
                seen_lh.add(lh.get("id")); all_lh_dedup.append(lh)

    # 3. CAPA DE DATOS
    build_data_layer_section(doc, folders, ws_lakehouses_map, all_dw)
    doc.add_page_break()

    # 4. CAPA DE ORQUESTACION (4.1, 4.2, 4.2.1, 4.2.1.1/2/3)
    build_orchestration_section(doc, data["workspaces"], folders, all_pipelines)

    # 4.3 ALMACENAMIENTO
    build_storage_section(doc, storage_folder, all_lh_dedup, all_dw, figure_n=8)

    # 4.4 CANALIZACION DE DESPLIEGUE
    build_deployment_section(doc, data.get("deployment_pipelines",[]), figure_n=9)

    build_reports_section(doc, all_reports)
    doc.add_page_break()

    # 5. SEMANTIC MODELS
    if all_sm or POWERBI_TABS:
        build_semantic_models_section(doc, all_sm)

    build_back_cover(doc)
    doc = build_static_toc(doc)
    return doc


print("OK Celda 6e v3.3 cargada")
