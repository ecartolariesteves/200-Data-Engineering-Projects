def _build_bronze_origin_text(doc):
    # Si no hay ninguna fuente configurada, no escribir nada
    if not BRONZE_SOURCES and not BRONZE_EXCEL_SOURCES:
        return

    origin_parts = []
    for src in BRONZE_SOURCES:
        system = src.get("system", "")
        conn   = src.get("conexion", "").upper()
        server = src.get("server") or ""
        db     = src.get("database") or ""
        if conn == "SERVER" and server:
            if db:
                origin_parts.append(
                    f"{system} (servidor: {server}, base de datos: {db})")
            else:
                origin_parts.append(f"{system} (servidor: {server})")
        else:
            origin_parts.append(system)

    if origin_parts:
        p  = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r1 = p.add_run("En nuestro caso, contamos con tablas que provienen de ")
        r1.font.name = "Arial"; r1.font.size = Pt(9)
        for i, part in enumerate(origin_parts):
            rb = p.add_run(part)
            rb.font.name = "Arial"; rb.font.size = Pt(9); rb.bold = True
            if i < len(origin_parts) - 1:
                p.add_run(", ").font.name = "Arial"
        r2 = p.add_run(".")
        r2.font.name = "Arial"; r2.font.size = Pt(9)

    if BRONZE_EXCEL_SOURCES:
        doc.add_paragraph()
        for es in BRONZE_EXCEL_SOURCES:
            origen = es.get("origen", "SharePoint")
            site   = es.get("site", "")
            url    = es.get("url", "")
            files  = es.get("files", [])
            p  = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            r1 = p.add_run(
                "Tambien se han cargado ficheros Excel que provienen del sitio ")
            r1.font.name = "Arial"; r1.font.size = Pt(9)
            rb = p.add_run(origen)
            rb.font.name = "Arial"; rb.font.size = Pt(9); rb.bold = True
            if site:
                rs = p.add_run(f" - {site}")
                rs.font.name = "Arial"; rs.font.size = Pt(9)
            # Enlace URL si disponible
            if url:
                p.add_run(" (").font.name = "Arial"
                from docx.oxml.ns import qn
                from docx.oxml   import OxmlElement
                import re
                r_link  = p.add_run(url)
                r_link.font.name  = "Arial"
                r_link.font.size  = Pt(9)
                r_link.font.color.rgb = RGBColor(0x00, 0x70, 0xC0)
                r_link.font.underline = True
                p.add_run(")").font.name = "Arial"
            p.add_run(":").font.name = "Arial"
            for fname in files:
                pf = doc.add_paragraph(style="List Bullet")
                pf.paragraph_format.space_after = Pt(2)
                pf.add_run(fname).font.name = "Arial"


def _lh_workspace_table(doc, lh_name, ws_lakehouses):
    rows = sorted([
        [ws_name, lh.get("id", "--"), lh.get("sqlEndpoint", "--") or "--"]
        for ws_name, lh in ws_lakehouses
    ], key=lambda r: r[0])
    schema_table(doc,
        ["Workspace", "ID", "SQL Endpoint"],
        rows, col_widths=[2000, 2800, 4400])


STAGING_SCHEMAS_SET = {
    "StagingLakehouseForDataflows", "DataflowsStagingLakehouse",
    "DataflowsStagingWarehouse",    "StagingWarehouseForDataflows"
}


def _get_real_tables(tables):
    return [
        t for t in tables
        if t.get("schema", "") not in STAGING_SCHEMAS_SET
        and not any(
            t.get("schema", "").startswith(p)
            for p in ("StagingLakehouseForDataflows", "DataflowsStagingLakehouse",
                      "DataflowsStagingWarehouse",    "StagingWarehouseForDataflows")
        )
    ]


def _build_bronze_lakehouse_content(doc, lh_name, ws_lh_list,
                                     origenes_str, excel_suffix):
    lh     = ws_lh_list[0][1]
    tables = lh.get("tables", [])
    files  = lh.get("files", [])
    real_t = _get_real_tables(tables)

    add_h3(doc, f"  Lakehouse: {lh_name}")
    _lh_workspace_table(doc, lh_name, ws_lh_list)

    add_h2(doc, "3.2.1 Tablas")
    add_body(doc, TEXT_BRONZE_TABLES.format(
        origenes_str=origenes_str, excel_suffix=excel_suffix))
    doc.add_paragraph()

    if real_t:
        bronze_t = sorted(
            [t for t in real_t
             if t.get("schema", "").lower() in ("bronze", "dbo", "")],
            key=lambda t: (t.get("schema", ""), t.get("name", ""))
        ) or sorted(real_t, key=lambda t: (t.get("schema",""), t.get("name","")))
        rows = [
            [f"{t.get('schema','')}.{t.get('name','')}"
             if t.get("schema") else t.get("name", ""),
             t.get("format", "delta"),
             str(len(t.get("schema_cols", [])))]
            for t in bronze_t
        ]
        schema_table(doc, ["Tabla", "Formato", "No columnas"], rows,
                     col_widths=[5000, 1600, 1600])
        _add_table_caption(doc, "Tabla 1. Listado de tablas de la capa Bronze")
    elif files:
        EXCLUDE_FOLDERS = {"docs", "config", "img", ".trident", "system"}
        visible_files = [
            f for f in files
            if f.get("name","").lower() not in EXCLUDE_FOLDERS
            and not f.get("name","").startswith(".")
        ]
        if visible_files:
            add_body(doc,
                "Los datos de esta capa se almacenan como ficheros en el "
                "area Files del Lakehouse, organizados en las siguientes carpetas:")
            doc.add_paragraph()
            rows = sorted([
                [f.get("name",""),
                 ", ".join(f.get("formats",[])) if f.get("formats")
                 else ("subcarpetas" if f.get("subfolders") else "ficheros sin extension"),
                 "\n".join(sorted(
                     f.get("files",[]) + f.get("subfolders",[])
                 )) or "--"]
                for f in visible_files
            ], key=lambda r: r[0])
            schema_table(doc, ["Carpeta", "Formato", "Contenido"], rows,
                         col_widths=[2400, 1800, 4000])
            _add_table_caption(doc,
                "Tabla 1. Estructura de carpetas de datos en el Lakehouse")
        else:
            add_body(doc, "Sin datos detectados en el Lakehouse.")
    else:
        add_body(doc, "Sin datos detectados en el Lakehouse.")


def build_data_layer_section(doc, folder_names, ws_lakehouses_map, all_warehouses):
    bronze_folder = folder_names.get("bronze", "001 - Data - Bronze")
    silver_folder = folder_names.get("silver", "002 - Data - Silver")
    gold_folder   = folder_names.get("gold",   "003 - Data - Gold")
    bold_folders  = [bronze_folder, silver_folder, gold_folder,
                     "Bronze", "Silver", "Gold"]
    has_excel     = bool(BRONZE_EXCEL_SOURCES)

    real_dw = [wh for wh in all_warehouses
               if not wh["name"].startswith("StagingWarehouseForDataflows")
               and not wh["name"].startswith("DataflowsStagingWarehouse")]

    dw_has_tables = any(
        len(t.get("schema_cols", [])) > 0
        for dw in real_dw
        for t in dw.get("tables", [])
    )

    bronze_dw = get_dw_tables_by_layer(real_dw, "bronze") if real_dw else []
    silver_dw = get_dw_tables_by_layer(real_dw, "silver") if real_dw else []
    gold_dw   = get_dw_tables_by_layer(real_dw, "gold")   if real_dw else []

    # 3. CAPA DE DATOS
    add_h1(doc, "3. Capa de datos")
    _add_body_with_bold(doc,
        TEXT_CAPA_DATOS_P1.format(
            bronze=bronze_folder, silver=silver_folder, gold=gold_folder),
        bold_folders)
    doc.add_paragraph()
    _add_body_with_bold(doc, TEXT_CAPA_DATOS_P2, bold_folders)
    doc.add_paragraph()
    _add_body_with_bold(doc,
        TEXT_CAPA_DATOS_P3.format(cliente=DOC_CLIENT), bold_folders)
    doc.add_paragraph()

    # 3.1 NOMENCLATURA
    add_h2(doc, "3.1 Nomenclatura")
    add_body(doc, TEXT_NOMENCLATURA)
    doc.add_paragraph()

    add_h2(doc, "3.1.1 Esquemas")
    add_body(doc, TEXT_ESQUEMAS)
    doc.add_paragraph()
    for key, desc in BULLETS_ESQUEMAS:
        _add_bullet_with_bold_key(doc, key, desc)
    doc.add_paragraph()

    add_h2(doc, "3.1.2 Tablas")
    add_body(doc, TEXT_TABLAS_INTRO)
    doc.add_paragraph()
    for key, desc in BULLETS_TABLAS_ESTRUCTURA:
        _add_bullet_with_bold_key(doc, key, desc)
    doc.add_paragraph()
    add_body(doc, TEXT_TABLAS_DIM)
    doc.add_paragraph()
    for key, desc in BULLETS_TABLAS_TIPO:
        _add_bullet_with_bold_key(doc, key, desc)
    doc.add_paragraph()

    # 3.2 BRONZE
    add_h1(doc, "3.2 Bronze")
    add_body(doc, TEXT_BRONZE_P1)
    doc.add_paragraph()
    add_body(doc, TEXT_BRONZE_P2)
    doc.add_paragraph()
    _build_bronze_origin_text(doc)
    doc.add_paragraph()

    origenes_str = ", ".join(
        s.get("system", "") for s in BRONZE_SOURCES if s.get("system")
    ) or "los sistemas de origen"
    excel_suffix = " y de los ficheros Excel" if has_excel else ""

    for lh_name, ws_lh_list in ws_lakehouses_map.items():
        _build_bronze_lakehouse_content(
            doc, lh_name, ws_lh_list, origenes_str, excel_suffix)

    if not ws_lakehouses_map:
        add_body(doc, "No se encontraron Lakehouses.")
    doc.add_paragraph()

    # 3.3 SILVER
    add_h1(doc, "3.3 Silver")
    add_body(doc, TEXT_SILVER_P1)
    doc.add_paragraph()
    add_body(doc, TEXT_SILVER_P2)
    doc.add_paragraph()

    add_h2(doc, "3.3.1 Tablas")
    seen_silver = set()
    silver_lh   = []
    for lh_name, ws_lh_list in ws_lakehouses_map.items():
        lh = ws_lh_list[0][1]
        for t in _get_real_tables(lh.get("tables", [])):
            if t.get("schema", "").lower() == "silver":
                key = t.get("name", "")
                if key not in seen_silver:
                    seen_silver.add(key)
                    silver_lh.append(t)
    silver_lh.sort(key=lambda t: t.get("name", ""))

    silver_dw_extra = [t for t in silver_dw if t.get("name", "") not in seen_silver]
    silver_all      = silver_lh + silver_dw_extra

    if silver_all:
        rows = [
            [f"{t.get('schema', 'silver')}.{t.get('name', '')}",
             t.get("format", "delta"),
             str(len(t.get("schema_cols", [])))]
            for t in silver_all
        ]
        schema_table(doc, ["Tabla", "Formato", "No columnas"], rows,
                     col_widths=[5000, 1600, 1600])
        _add_table_caption(doc, "Tabla 2. Listado de tablas de la capa Silver")
        if silver_dw_extra and not silver_lh:
            dw_schemas = sorted(set(t.get("schema", "") for t in silver_dw_extra))
            p  = doc.add_paragraph()
            r1 = p.add_run(
                f"Las tablas de esta capa se encuentran en el Data Warehouse "
                f"({real_dw[0]['name']}) bajo los schemas: "
                f"{', '.join(dw_schemas)}.")
            r1.font.name = "Arial"; r1.font.size = Pt(9)
        elif silver_dw_extra and silver_lh:
            p  = doc.add_paragraph()
            r1 = p.add_run(
                f"Tablas del Lakehouse: {len(silver_lh)}. "
                f"Tablas del Data Warehouse: {len(silver_dw_extra)}.")
            r1.font.name = "Arial"; r1.font.size = Pt(9)
    elif real_dw and not dw_has_tables:
        p  = doc.add_paragraph()
        r1 = p.add_run(
            f"Las tablas de esta capa se encuentran en el Data Warehouse "
            f"({real_dw[0]['name']}). ")
        r1.font.name = "Arial"; r1.font.size = Pt(9)
        r2 = p.add_run(
            "El esquema detallado esta disponible en el SQL Analytics Endpoint.")
        r2.font.name = "Arial"; r2.font.size = Pt(9); r2.font.italic = True
    else:
        add_body(doc, "No se detectaron tablas en la capa Silver.")
    doc.add_paragraph()

    # 3.4 GOLD
    add_h1(doc, "3.4 Gold")
    add_body(doc, TEXT_GOLD_P1)
    doc.add_paragraph()
    for bullet in BULLETS_GOLD_OPS:
        _add_bullet_plain(doc, bullet)
    doc.add_paragraph()

    # Recoger tablas gold del LH
    seen_gold = set()
    gold_lh   = []
    for lh_name, ws_lh_list in ws_lakehouses_map.items():
        lh = ws_lh_list[0][1]
        for t in _get_real_tables(lh.get("tables", [])):
            if t.get("schema", "").lower() == "gold":
                key = t.get("name", "")
                if key not in seen_gold:
                    seen_gold.add(key)
                    gold_lh.append(t)
    gold_lh.sort(key=lambda t: t.get("name", ""))

    gold_dw_extra = [t for t in gold_dw if t.get("name", "") not in seen_gold]
    gold_all      = gold_lh + gold_dw_extra

    # Separar tablas fisicas de vistas
    gold_tables_only = [t for t in gold_all if t.get("type", "TABLE") != "VIEW"]
    gold_views_only  = [t for t in gold_all if t.get("type", "TABLE") == "VIEW"]

    td_tables = [t for t in gold_tables_only if "_TD_" in t.get("name", "").upper()]
    th_tables = [t for t in gold_tables_only if "_TH_" in t.get("name", "").upper()]
    other_t   = [t for t in gold_tables_only
                 if "_TD_" not in t.get("name", "").upper()
                 and "_TH_" not in t.get("name", "").upper()]

    # 3.4.1 TABLAS
    add_h2(doc, "3.4.1 Tablas")
    if gold_tables_only:
        n_td = len(td_tables); n_th = len(th_tables)
        gold_desc = (
            f"El modelo Gold esta compuesto por {n_td} tablas de dimension "
            f"y {n_th} tablas de hechos. "
            "Las tablas de dimension (prefijo TD_) contienen los datos maestros "
            "y de referencia del modelo. "
            "Las tablas de hechos (prefijo TH_) registran los eventos y "
            "transacciones del negocio."
        )
        _add_body_with_bold(doc, gold_desc,
                            ["tablas de dimension", "tablas de hechos", "TD_", "TH_"])
        doc.add_paragraph()
        for label, tbls in [("Dimensiones:", td_tables),
                             ("Hechos:",      th_tables),
                             ("Otras tablas:", other_t)]:
            if tbls:
                lbl = doc.add_paragraph()
                lbl.add_run(label).bold = True
                rows = [
                    [f"{t.get('schema', 'gold')}.{t.get('name', '')}",
                     t.get("format", "delta"),
                     str(len(t.get("schema_cols", [])))]
                    for t in tbls
                ]
                schema_table(doc, ["Tabla", "Formato", "No columnas"], rows,
                             col_widths=[5000, 1600, 1600])
        _add_table_caption(doc, "Tabla 3. Listado de tablas de la capa Gold")
        if gold_dw_extra and not gold_lh:
            dw_schemas = sorted(set(t.get("schema", "") for t in gold_dw_extra
                                    if t.get("type", "TABLE") != "VIEW"))
            if dw_schemas:
                p  = doc.add_paragraph()
                r1 = p.add_run(
                    f"Las tablas de esta capa se encuentran en el Data Warehouse "
                    f"({real_dw[0]['name']}) bajo los schemas: "
                    f"{', '.join(dw_schemas)}.")
                r1.font.name = "Arial"; r1.font.size = Pt(9)
    elif real_dw and not dw_has_tables:
        p  = doc.add_paragraph()
        r1 = p.add_run(
            f"Las tablas de esta capa se encuentran en el Data Warehouse "
            f"({real_dw[0]['name']}). ")
        r1.font.name = "Arial"; r1.font.size = Pt(9)
        r2 = p.add_run(
            "El esquema detallado esta disponible en el SQL Analytics Endpoint.")
        r2.font.name = "Arial"; r2.font.size = Pt(9); r2.font.italic = True
    else:
        add_body(doc, "No se detectaron tablas en la capa Gold.")
    doc.add_paragraph()

    # 3.4.2 VISTAS
    add_h2(doc, "3.4.2 Vistas")
    add_body(doc,
        "Las vistas del modelo Gold actuan como capa de acceso optimizada para "
        "las herramientas de visualizacion y consumo de datos. A diferencia de "
        "las tablas fisicas, las vistas no almacenan datos propios sino que "
        "exponen una proyeccion estructurada y controlada del modelo dimensional "
        "subyacente.")
    doc.add_paragraph()
    add_body(doc,
        "Su funcion principal es simplificar el acceso desde herramientas como "
        "Power BI, permitiendo que los modelos semanticos consuman datos ya "
        "preparados, con las relaciones, filtros y transformaciones aplicadas, "
        "sin necesidad de logica adicional en la capa de visualizacion. El acceso "
        "se realiza a traves del SQL Analytics Endpoint del Data Warehouse.")
    doc.add_paragraph()
    if gold_views_only:
        rows = [
            [f"{t.get('schema', 'gold')}.{t.get('name', '')}",
             str(len(t.get("schema_cols", [])))]
            for t in sorted(gold_views_only, key=lambda t: t.get("name", ""))
        ]
        schema_table(doc, ["Vista", "No columnas"], rows,
                     col_widths=[6200, 1600])
        _add_table_caption(doc, "Tabla 4. Listado de vistas de la capa Gold")
    else:
        add_body(doc, "No se detectaron vistas en la capa Gold.")
    doc.add_paragraph()

    # 3.4.3 DATA WAREHOUSES
    add_h2(doc, "3.4.3 Data Warehouses")
    if real_dw:
        for wh in real_dw:
            add_h3(doc, f"  DW: {wh.get('name','')}")
            meta_table(doc, {
                "ID":                wh.get("id", "--"),
                "Connection String": wh.get("connectionString", "") or "--",
            })
            if dw_has_tables:
                wh_tables     = wh.get("tables", [])
                schemas_found = sorted(set(
                    t.get("schema", "") for t in wh_tables
                    if t.get("schema", "") not in ("", "(esquema no accesible)")
                ))
                if schemas_found:
                    p  = doc.add_paragraph()
                    r1 = p.add_run(
                        f"Schemas detectados: {', '.join(schemas_found)}. "
                        f"Total objetos: {len(wh_tables)}.")
                    r1.font.name = "Arial"; r1.font.size = Pt(9)
            else:
                note = doc.add_paragraph()
                rn   = note.add_run(
                    "El esquema del Data Warehouse no es accesible via Spark. "
                    "Consultar directamente el SQL Analytics Endpoint.")
                rn.font.name = "Arial"; rn.font.size = Pt(8)
                rn.font.italic = True; rn.font.color.rgb = GRAY_MID
        doc.add_paragraph()
    else:
        add_body(doc, "No se detectaron Data Warehouses en el workspace.")
        doc.add_paragraph()


print("OK CELL_6d_v33 cargada -- soporte LH Delta + LH ficheros + DW + Vistas Gold")
