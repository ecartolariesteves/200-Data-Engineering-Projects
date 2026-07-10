import os as _os
import io
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.table import WD_TABLE_ALIGNMENT

# ── Paleta FabricDoc ─────────────────────────────────────────────────────────
ACCENT       = RGBColor(0x00, 0x78, 0xD4)  # Microsoft Fabric blue
GRAY_DARK    = RGBColor(0x40, 0x40, 0x40)
GRAY_MID     = RGBColor(0x80, 0x80, 0x80)
WHITE        = RGBColor(0xFF, 0xFF, 0xFF)
TABLE_HDR_BG = "0078D4"
TABLE_ALT_BG = "F0F6FF"
BORDER_COLOR = "E0E0E0"

def set_cell_bg(cell, hex_color):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_color)
    tcPr.append(shd)


def set_cell_borders(cell, color=BORDER_COLOR):
    tc       = cell._tc
    tcPr     = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for side in ["top", "left", "bottom", "right"]:
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"),   "single")
        el.set(qn("w:sz"),    "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def style_cell(cell, text, bold=False, color=None, bg=None,
               font_size=9, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p   = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(str(text) if text is not None else "")
    run.bold      = bold
    run.font.size = Pt(font_size)
    run.font.name = "Arial"
    if color:
        run.font.color.rgb = color
    if bg:
        set_cell_bg(cell, bg)
    set_cell_borders(cell)
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for side, val in [("top", 80), ("bottom", 80), ("left", 120), ("right", 120)]:
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"),    str(val))
        el.set(qn("w:type"), "dxa")
        tcMar.append(el)
    tcPr.append(tcMar)


def add_h1(doc, text):
    p   = doc.add_paragraph()
    p.style = doc.styles["Heading 1"]
    run = p.add_run(text.upper())
    run.font.name      = "Arial"
    run.font.size      = Pt(13)
    run.font.bold      = True
    run.font.color.rgb = GRAY_DARK
    pPr    = p._p.get_or_add_pPr()
    pBdr   = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"),   "single")
    bottom.set(qn("w:sz"),    "12")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "0078D4")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def add_h2(doc, text):
    p   = doc.add_paragraph()
    p.style = doc.styles["Heading 2"]
    run = p.add_run(text)
    run.font.name       = "Arial"
    run.font.size       = Pt(10)
    run.font.bold       = True
    run.font.color.rgb  = ACCENT
    run.font.small_caps = True
    return p


def add_h3(doc, text):
    p   = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after  = Pt(2)
    run = p.add_run(text)
    run.font.name      = "Arial"
    run.font.size      = Pt(9)
    run.font.bold      = True
    run.font.italic    = True
    run.font.color.rgb = GRAY_DARK
    return p


def add_body(doc, text):
    p   = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(4)    # solo 4pt entre párrafos de texto
    p.paragraph_format.line_spacing = Pt(13)
    run = p.add_run(str(text))
    run.font.name = "Arial"
    run.font.size = Pt(9)
    return p


def meta_table(doc, items, col_widths=(2.0, 4.7)):
    table = doc.add_table(rows=0, cols=2)
    table.style     = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    hdr = table.add_row()
    style_cell(hdr.cells[0], "Propiedad", bold=True, color=WHITE, bg=TABLE_HDR_BG)
    style_cell(hdr.cells[1], "Valor",     bold=True, color=WHITE, bg=TABLE_HDR_BG)
    for i, (k, v) in enumerate(items.items()):
        row = table.add_row()
        bg  = TABLE_ALT_BG if i % 2 == 1 else "FFFFFF"
        style_cell(row.cells[0], k,        bold=True, bg=bg)
        style_cell(row.cells[1], v or "—", bg=bg)
    # Párrafo separador compacto
    sep = doc.add_paragraph()
    sep.paragraph_format.space_before = Pt(0)
    sep.paragraph_format.space_after  = Pt(6)


def schema_table(doc, headers, rows, col_widths=None):
    n = len(headers)
    if not col_widths:
        total      = int(6.7 * 1440)
        col_widths = [total // n] * n
    table = doc.add_table(rows=0, cols=n)
    table.style = "Table Grid"
    hdr = table.add_row()
    for i, h in enumerate(headers):
        style_cell(hdr.cells[i], h, bold=True, color=WHITE,
                   bg=TABLE_HDR_BG, font_size=8.5)
    for ri, row_data in enumerate(rows):
        row = table.add_row()
        bg  = TABLE_ALT_BG if ri % 2 == 1 else "FFFFFF"
        for ci, val in enumerate(row_data):
            style_cell(row.cells[ci], val, font_size=8.5, bg=bg)
    # Párrafo separador compacto
    sep = doc.add_paragraph()
    sep.paragraph_format.space_before = Pt(0)
    sep.paragraph_format.space_after  = Pt(6)


def clean_comment(val):
    if val is None or str(val).strip().lower() in ("none", "null", ""):
        return ""
    return str(val).strip()


def _add_table_caption(doc, caption_text):
    p   = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(caption_text)
    run.font.name      = "Arial"
    run.font.size      = Pt(8)
    run.font.italic    = True
    run.font.bold      = True
    run.font.color.rgb = GRAY_MID


def _add_body_with_bold(doc, text, bold_words):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(4)
    p.paragraph_format.line_spacing = Pt(13)
    remaining = text
    while remaining:
        first_pos  = len(remaining)
        first_word = None
        for bw in sorted(bold_words, key=len, reverse=True):
            idx = remaining.find(bw)
            if idx != -1 and idx < first_pos:
                first_pos  = idx
                first_word = bw
        if first_word is None:
            run = p.add_run(remaining)
            run.font.name = "Arial"; run.font.size = Pt(9)
            break
        if first_pos > 0:
            run = p.add_run(remaining[:first_pos])
            run.font.name = "Arial"; run.font.size = Pt(9)
        run = p.add_run(first_word)
        run.font.name = "Arial"; run.font.size = Pt(9); run.font.bold = True
        remaining = remaining[first_pos + len(first_word):]
    return p


def _add_bullet_plain(doc, text):
    p   = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(2)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(9)
    return p


def _add_bullet_with_bold_key(doc, key, description):
    p       = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(2)
    run_key = p.add_run(key)
    run_key.font.name = "Arial"; run_key.font.size = Pt(9); run_key.bold = True
    run_sep = p.add_run(f": {description}")
    run_sep.font.name = "Arial"; run_sep.font.size = Pt(9)
    return p


print("✅ Celda 6a cargados")